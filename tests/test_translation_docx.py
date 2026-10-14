"""A translation's HTML as a Word document.

Every case builds a real .docx and reads it back with python-docx, so what is
checked is the file a client would open, not the converter's own bookkeeping.
The HTML is what Suliko Translate writes (headings, paragraphs, lists, tables,
colour spans) plus what the page that corrects it adds (`div`, `br`, `b`, `i`).
"""

from __future__ import annotations

import contextlib
import io
import zipfile
from typing import Any

import pytest
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from suliko.domain import translation_docx
from suliko.domain.translation_docx import DocxConversionError, html_to_docx

#: A one-pixel PNG.
PNG = "".join(
    [
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA",
        "60e6kgAAAABJRU5ErkJggg==",
    ]
)


def convert(html: str, **kwargs: Any) -> DocumentObject:
    return Document(io.BytesIO(html_to_docx(html, **kwargs)))


def texts(document: DocumentObject) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs]


# ── The document ────────────────────────────────────────────────────────────


def test_it_is_a_real_word_file_on_an_a4_page_in_calibri() -> None:
    raw = html_to_docx("<p>Hello</p>", language="en")
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = set(archive.namelist())
        assert "word/document.xml" in names and "[Content_Types].xml" in names
        # Not the HTML-in-a-document wrapper suliko.ge's own export writes.
        assert not any(name.endswith((".html", ".htm")) for name in names)
        assert b"altChunk" not in archive.read("word/document.xml")

    document = Document(io.BytesIO(raw))
    section = document.sections[0]
    assert round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7
    assert round(section.left_margin.cm, 1) == 2.0 and round(section.top_margin.cm, 1) == 2.0
    normal = document.styles["Normal"]
    assert normal.font.name == "Calibri" and normal.font.size is not None
    assert normal.font.size.pt == 11
    assert document.core_properties.author == "Suliko Translate"


def test_the_language_of_the_translation_is_what_word_spell_checks() -> None:
    document = convert("<p>x</p>", language="ka")
    lang = document.styles["Normal"].element.rPr.find(qn("w:lang"))
    assert lang is not None and lang.get(qn("w:val")) == "ka"
    # Anything that is not a language code is not written into the file.
    document = convert("<p>x</p>", language='"><x/>')
    assert document.styles["Normal"].element.rPr.find(qn("w:lang")) is None


def test_the_title_becomes_the_documents_title() -> None:
    document = convert("<html><head><title>Passport</title></head><body><p>x</p></body></html>")
    assert document.core_properties.title == "Passport"


def test_an_empty_translation_is_still_a_document() -> None:
    for html in ("", "   \n", "<script>1</script>", "<!-- nothing -->"):
        document = convert(html)
        assert [p.text for p in document.paragraphs if p.text] == []


def test_georgian_survives_in_every_place_it_can_be() -> None:
    document = convert(
        "<h1>საქართველოს პასპორტი</h1><p>გამარჯობა, <strong>ნინო</strong></p>"
        "<table><tr><td>თარიღი</td></tr></table><ul><li>პუნქტი</li></ul>"
    )
    assert "საქართველოს პასპორტი" in texts(document)
    assert "გამარჯობა, ნინო" in texts(document)
    assert document.tables[0].cell(0, 0).text == "თარიღი"
    assert "პუნქტი" in texts(document)


# ── Paragraphs and text ─────────────────────────────────────────────────────


def test_headings_use_the_headings_styles_in_black() -> None:
    document = convert("<h1>One</h1><h2>Two</h2><h6>Six</h6><p>Body</p>")
    assert [(p.text, p.style.name) for p in document.paragraphs] == [
        ("One", "Heading 1"),
        ("Two", "Heading 2"),
        ("Six", "Heading 6"),
        ("Body", "Normal"),
    ]
    heading = document.styles["Heading 1"]
    assert heading.font.name == "Calibri" and heading.font.bold is True
    assert str(heading.font.color.rgb) == "000000"


def test_whitespace_is_collapsed_the_way_a_browser_does() -> None:
    document = convert("<p>  Last   paragraph\n  with\tspaces  </p><p>&nbsp;kept&nbsp;</p>")
    assert texts(document) == ["Last paragraph with spaces", "\u00a0kept\u00a0"]


def test_a_line_break_ends_a_line_and_one_at_the_end_of_a_block_is_nothing() -> None:
    document = convert("<p>a<br>b</p><p>c<br></p><p>d<br><br>e</p><p><br></p>")
    assert texts(document) == ["a\nb", "c", "d\n\ne", ""]


def test_text_between_blocks_gets_a_paragraph_and_nothing_empty_is_made() -> None:
    document = convert("<div>Bare<div>inner</div>tail</div>")
    assert texts(document) == ["Bare", "inner", "tail"]


def test_emphasis_colour_size_and_shading_reach_the_runs() -> None:
    document = convert(
        "<p>plain <strong>bold</strong> <em>italic</em> <u>under</u> <s>struck</s> "
        'x<sup>2</sup> H<sub>2</sub>O <span style="color:#C00000">red</span> '
        '<span style="font-size:16px;background-color:#ffff00">big</span> '
        '<font color="blue" size="5">font</font> '
        '<span style="font-weight:700">heavy</span></p>'
    )
    by_text: dict[str, list[Any]] = {}
    for run in document.paragraphs[0].runs:
        if run.text.strip():
            by_text.setdefault(run.text.strip(), []).append(run)
    runs = {text: found[0] for text, found in by_text.items()}
    assert runs["bold"].bold and runs["italic"].italic and runs["under"].underline
    assert runs["struck"].font.strike
    superscript, subscript = by_text["2"]
    assert superscript.font.superscript is True and subscript.font.subscript is True
    assert str(runs["red"].font.color.rgb) == "C00000"
    big = runs["big"]
    assert big.font.size is not None and big.font.size.pt == 12  # 16px
    assert big._r.xpath("./w:rPr/w:shd/@w:fill") == ["FFFF00"]
    assert str(runs["font"].font.color.rgb) == "0000FF"
    assert runs["font"].font.size is not None and runs["font"].font.size.pt == 18
    assert runs["heavy"].bold is True
    assert runs["plain"].bold is None


def test_the_space_before_a_decorated_word_is_not_decorated() -> None:
    document = convert('<p>a <u>b</u> c <span style="background-color:#ff0">d</span></p>')
    runs = document.paragraphs[0].runs
    assert "".join(run.text for run in runs) == "a b c d"
    for run in runs:
        if run.text.strip() == "":
            assert not run.underline and not run._r.xpath("./w:rPr/w:shd")
        if run.text.strip() in {"b", "d"}:
            assert run.text == run.text.strip()  # the space is in a run of its own


def test_alignment_and_indent_from_style_or_attribute() -> None:
    document = convert(
        '<p style="text-align:center">c</p><p align="right">r</p>'
        '<p style="text-align:justify;margin-left:20px">j</p><blockquote>q</blockquote>'
    )
    paragraphs = document.paragraphs
    assert paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraphs[2].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraphs[2].paragraph_format.left_indent is not None
    assert paragraphs[2].paragraph_format.left_indent.pt == 15  # 20px
    assert paragraphs[3].paragraph_format.left_indent is not None
    assert paragraphs[3].paragraph_format.left_indent.pt == 36


def test_a_rule_and_a_page_break() -> None:
    document = convert(
        '<p>one</p><hr><p>two</p><div style="page-break-after: always"></div><p>three</p>'
        '<p style="break-before: page">four</p>'
    )
    breaks = [p.paragraph_format.page_break_before for p in document.paragraphs if p.text]
    assert breaks == [None, None, True, True]
    assert document.paragraphs[1]._p.xpath("./w:pPr/w:pBdr/w:bottom")


def test_preformatted_text_keeps_its_lines_and_spaces() -> None:
    document = convert("<pre>a\n  b</pre>")
    assert texts(document) == ["a\n  b"]
    assert document.paragraphs[0].runs[0].font.name == "Consolas"


# ── Links and images ────────────────────────────────────────────────────────


def test_only_a_safe_address_becomes_a_link() -> None:
    document = convert(
        '<p><a href="https://suliko.ge/ka">site</a> <a href="mailto:a@b.ge">mail</a> '
        '<a href="javascript:alert(1)">bad</a> <a href="file:///c:/x">local</a></p>'
    )
    paragraph = document.paragraphs[0]
    links = paragraph._p.xpath("./w:hyperlink")
    assert len(links) == 2
    targets = {document.part.rels[link.get(qn("r:id"))].target_ref for link in links}
    assert targets == {"https://suliko.ge/ka", "mailto:a@b.ge"}
    # The words of the refused ones stay.
    assert paragraph.text == "site mail bad local"


def test_an_image_inside_the_html_is_kept_and_one_to_fetch_is_not() -> None:
    document = convert(
        f'<p>x<img src="data:image/png;base64,{PNG}" width="40"></p>'
        '<p>y<img src="https://example.com/logo.png"><img src="file:///c:/secret.png"></p>'
    )
    assert len(document.inline_shapes) == 1
    assert texts(document) == ["x", "y"]
    # A broken image is left out and the translation is not.
    assert texts(convert('<p>z<img src="data:image/png;base64,AAAA"></p>')) == ["z"]


# ── Lists ───────────────────────────────────────────────────────────────────


def _numbering(paragraph: Any) -> tuple[int | None, int | None]:
    level = paragraph._p.xpath("./w:pPr/w:numPr/w:ilvl/@w:val")
    number = paragraph._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")
    return (int(number[0]) if number else None, int(level[0]) if level else None)


def test_lists_nest_by_level_and_bullets_share_one_numbering() -> None:
    document = convert(
        "<ul><li>a</li><li>b<ul><li>b1</li><li>b2</li></ul></li><li>c</li></ul>"
        "<p>between</p><ul><li>d</li></ul>"
    )
    paragraphs = {p.text: _numbering(p) for p in document.paragraphs}
    bullet = paragraphs["a"][0]
    assert bullet is not None
    assert paragraphs == {
        "a": (bullet, 0),
        "b": (bullet, 0),
        "b1": (bullet, 1),
        "b2": (bullet, 1),
        "c": (bullet, 0),
        "between": (None, None),
        "d": (bullet, 0),
    }


def test_each_numbered_list_starts_again_and_honours_start() -> None:
    document = convert(
        "<ol><li>one</li><li>two<ol><li>two a</li></ol></li></ol><p>x</p>"
        '<ol><li>again</li></ol><ol start="5"><li>five</li></ol>'
    )
    paragraphs = {p.text: _numbering(p) for p in document.paragraphs}
    first = paragraphs["one"][0]
    assert first is not None
    assert paragraphs["two"] == (first, 0)
    assert paragraphs["again"][0] not in (None, first)
    assert paragraphs["five"][0] not in (None, first, paragraphs["again"][0])

    numbering = document.part.numbering_part.element
    overrides = {
        int(num.get(qn("w:numId"))): num.xpath("./w:lvlOverride/w:startOverride/@w:val")
        for num in numbering.findall(qn("w:num"))
    }
    assert overrides[first][0] == "1" and overrides[paragraphs["five"][0]][0] == "5"
    # Word wants every abstract numbering before the first numbering.
    tags = [child.tag.rsplit("}", 1)[1] for child in numbering]
    assert tags.index("num") > max(i for i, tag in enumerate(tags) if tag == "abstractNum")


def test_list_items_have_no_space_after_but_the_list_does() -> None:
    document = convert("<ul><li>a</li><li>b</li></ul><p>after</p>")
    first, second, after = document.paragraphs
    assert (
        first.paragraph_format.space_after is not None
        and first.paragraph_format.space_after.pt == 0
    )
    assert (
        second.paragraph_format.space_after is not None
        and second.paragraph_format.space_after.pt == 4
    )
    assert after.paragraph_format.space_after is None  # the body's own


# ── Tables ──────────────────────────────────────────────────────────────────


def test_a_table_keeps_its_cells_headers_spans_and_shading() -> None:
    document = convert(
        '<table border="1"><thead><tr><th>A</th><th>B</th><th>C</th></tr></thead><tbody>'
        '<tr><td colspan="2">wide</td>'
        '<td rowspan="2" style="background-color:#ddeeff">tall</td></tr>'
        "<tr><td>x</td><td>y</td></tr></tbody></table><p>after</p>"
    )
    [table] = document.tables
    assert (len(table.rows), len(table.columns)) == (3, 3)
    assert table.style is not None and table.style.name == "Table Grid"
    assert [c.text for c in table.rows[0].cells] == ["A", "B", "C"]
    header = table.cell(0, 0).paragraphs[0]
    assert header.runs[0].bold and header.alignment == WD_ALIGN_PARAGRAPH.CENTER
    wide = table.cell(1, 0)
    assert wide.text == "wide" and wide._tc.xpath("./w:tcPr/w:gridSpan/@w:val") == ["2"]
    tall = table.cell(1, 2)
    assert tall.text == "tall" and tall._tc.xpath("./w:tcPr/w:vMerge/@w:val") == ["restart"]
    assert tall._tc.xpath("./w:tcPr/w:shd/@w:fill") == ["DDEEFF"]
    assert [c.text for c in table.rows[2].cells[:2]] == ["x", "y"]
    # No space inside a cell, and some before the paragraph that follows.
    assert table.cell(1, 0).paragraphs[0].paragraph_format.space_after.pt == 0
    assert document.paragraphs[-1].text == "after"
    assert document.paragraphs[-1].paragraph_format.space_before.pt == 6


def test_a_table_is_drawn_with_lines_only_when_the_html_draws_them() -> None:

    def grid(html: str) -> bool:
        style = convert(html).tables[0].style
        return style is not None and style.name == "Table Grid"

    assert not grid("<table><tr><td>layout</td></tr></table>")
    assert not grid('<table border="0"><tr><td>x</td></tr></table>')
    assert not grid('<table><tr><td style="border:none">x</td></tr></table>')
    assert not grid('<table style="border-collapse:collapse"><tr><td>x</td></tr></table>')
    assert grid('<table border="1"><tr><td>x</td></tr></table>')
    assert grid('<table><tr><td style="border:1px solid #000">x</td></tr></table>')


def test_a_cell_can_hold_paragraphs_lists_and_another_table() -> None:
    document = convert(
        "<table><tr><td><p>one</p><p>two</p><ul><li>item</li></ul>"
        "<table><tr><td>inner</td></tr></table></td><td>next</td></tr></table>"
    )
    cell = document.tables[0].cell(0, 0)
    assert [p.text for p in cell.paragraphs if p.text] == ["one", "two", "item"]
    assert cell.tables[0].cell(0, 0).text == "inner"
    assert document.tables[0].cell(0, 1).text == "next"


def test_rows_of_different_lengths_and_crossing_spans_still_make_a_table() -> None:
    document = convert(
        '<table><tr><td rowspan="2">a</td><td>b</td></tr><tr><td colspan="3">c</td></tr>'
        "<tr><td>d</td></tr></table>"
    )
    table = document.tables[0]
    assert {cell.text for row in table.rows for cell in row.cells} >= {"a", "b", "c", "d"}


def test_a_caption_comes_before_its_table() -> None:
    document = convert("<table><caption>Dates</caption><tr><td>x</td></tr></table>")
    assert texts(document)[0] == "Dates"
    assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


# ── What it does not trust ──────────────────────────────────────────────────


def test_scripts_styles_and_anything_that_runs_or_fetches_are_left_out() -> None:
    document = convert(
        "<style>p{color:red}</style><script>alert(1)</script><p>kept</p>"
        "<iframe src='https://x.example/'>no</iframe><object data='x'>no</object>"
        '<input value="no"><button>no</button><!-- no -->'
        '<p onclick="evil()" class="x">still kept</p>'
    )
    assert texts(document) == ["kept", "still kept"]


def test_characters_xml_cannot_carry_are_dropped_not_fatal() -> None:
    document = convert("<p>a\x00b\x08c\x1fd\ufffee</p>")
    assert texts(document) == ["abcde"]


def test_unknown_and_broken_html_is_read_as_the_text_it_holds() -> None:
    document = convert(
        "<p>one <blink>two</blink> <custom-tag>three</custom-tag><p>four</i></b></div>"
    )
    assert texts(document) == ["one two three", "four"]
    assert texts(convert("plain text, no tags")) == ["plain text, no tags"]
    assert texts(convert("<td>cell outside a table</td>")) == ["cell outside a table"]
    assert texts(convert("<li>item outside a list</li>")) == ["item outside a list"]


def test_a_document_with_an_xml_prolog_or_byte_order_mark_is_read() -> None:
    assert texts(convert('\ufeff<?xml version="1.0" encoding="utf-8"?><p>x</p>')) == ["x"]


def test_too_deeply_nested_is_refused_not_crashed() -> None:
    with pytest.raises(DocxConversionError):
        html_to_docx("<div>" * 200 + "x" + "</div>" * 200)


def test_a_table_that_is_too_large_is_refused() -> None:
    cells = "<td>x</td>" * 64
    with pytest.raises(DocxConversionError):
        html_to_docx(f"<table><tr>{cells}</tr></table>")
    wide = "<tr>" + "<td>x</td>" * 63 + "</tr>"
    with pytest.raises(DocxConversionError):
        html_to_docx("<table>" + wide * 400 + "</table>")


def test_a_document_with_too_many_elements_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(translation_docx, "MAX_NODES", 50)
    with pytest.raises(DocxConversionError):
        html_to_docx("<p>x</p>" * 100)


def test_nothing_else_escapes_as_an_error() -> None:
    """Whatever the HTML, the answer is a document or a `DocxConversionError`."""
    samples = [
        "<table><tr></tr></table>",
        "<table></table>",
        "<ul></ul><ol></ol>",
        "<tr><td>x",
        '<td colspan="x" rowspan="-3">x</td>',
        '<table><tr><td colspan="9999">x</td></tr></table>',
        '<p style="color:;font-size:;;:;margin-left:abc">x</p>',
        "<p style='font-size:99999pt'>big</p><p style='font-size:0'>small</p>",
        "<a>no href</a><a href=''>empty</a>",
        "<img><img src=''><img src='data:'>",
        "<br><br><br>",
        "<hr><hr>",
        "<dl><dt>t</dt><dd>d</dd></dl>",
        "<p>" + "word " * 50_000 + "</p>",
    ]
    for html in samples:
        with contextlib.suppress(DocxConversionError):
            html_to_docx(html)


def test_the_same_html_gives_the_same_words_twice() -> None:
    html = "<h1>A</h1><p>b <em>c</em></p><table><tr><td>d</td></tr></table><ol><li>e</li></ol>"
    first, second = convert(html), convert(html)
    assert texts(first) == texts(second)
    assert [c.text for c in first.tables[0].rows[0].cells] == ["d"]
