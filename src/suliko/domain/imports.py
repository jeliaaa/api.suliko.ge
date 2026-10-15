"""Importing a bureau's own spreadsheet: clients, translators, notaries and orders.

Pure code, no database: reading a file, guessing which column is which, cleaning
what is typed in a cell, and judging each row. The endpoints in
``api/v1/imports.py`` feed it what already exists and save what it accepts.

A row is one of three things:

- ``ok``: will be created (it may carry warnings, such as an email that was not
  valid and was left out).
- ``duplicate``: already in the bureau, or earlier in the same file, matched by
  phone, email or ID number. Skipped, with the existing record named.
- ``error``: cannot be created, with the reason. Skipped.

Nothing here raises for a bad row: a messy sheet is the normal case.
"""

# ruff: noqa: RUF001  (Georgian, Russian and typographic characters are what people type)

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from collections.abc import Callable, Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Literal

from suliko.core.errors import ValidationError
from suliko.domain.cities import CITIES

#: The size of file and number of rows one import takes. The app sits behind a
#: host that refuses bodies above about 4.5 MB, so the byte limit is below it.
MAX_BYTES = 4 * 1024 * 1024
MAX_ROWS = 5000

Kind = Literal["clients", "translators", "notaries", "orders"]
KINDS: tuple[Kind, ...] = ("clients", "translators", "notaries", "orders")


# ── The fields of each kind ─────────────────────────────────────────────────


@dataclass(frozen=True)
class FieldSpec:
    key: str
    required: bool = False


FIELDS: dict[str, tuple[FieldSpec, ...]] = {
    "clients": (
        FieldSpec("name", required=True),
        FieldSpec("client_type"),
        FieldSpec("phone"),
        FieldSpec("email"),
        FieldSpec("personal_number"),
        FieldSpec("address"),
        FieldSpec("acquisition_source"),
        FieldSpec("notes"),
    ),
    "translators": (
        FieldSpec("name", required=True),
        FieldSpec("phone"),
        FieldSpec("email"),
        FieldSpec("office_address"),
        FieldSpec("comment"),
        FieldSpec("default_rate"),
        FieldSpec("bank_iban"),
        FieldSpec("can_notarize"),
        FieldSpec("works_oral"),
        FieldSpec("cities"),
        FieldSpec("language_pairs"),
    ),
}

#: Headers people use, in Georgian, English and Russian, by field. Compared after
#: ``norm_header``, so case, spaces and punctuation do not matter.
SYNONYMS: dict[str, dict[str, tuple[str, ...]]] = {
    "clients": {
        "name": (
            "name",
            "full name",
            "fullname",
            "client",
            "client name",
            "customer",
            "სახელი",
            "სახელი გვარი",
            "გვარი სახელი",
            "კლიენტი",
            "კლიენტის სახელი",
            "სახელი და გვარი",
            "имя",
            "фио",
            "клиент",
            "заказчик",
        ),
        "client_type": ("type", "client type", "ტიპი", "კლიენტის ტიპი", "тип"),
        "phone": (
            "phone",
            "tel",
            "telephone",
            "mobile",
            "phone number",
            "ტელეფონი",
            "ტელ",
            "მობილური",
            "ტელეფონის ნომერი",
            "телефон",
            "тел",
        ),
        "email": ("email", "e-mail", "mail", "ელფოსტა", "ელ ფოსტა", "იმეილი", "почта"),
        "personal_number": (
            "id",
            "id number",
            "personal number",
            "personal id",
            "tax id",
            "identification code",
            "პირადი ნომერი",
            "პირადობის ნომერი",
            "საიდენტიფიკაციო კოდი",
            "ს კ",
            "პ ნ",
        ),
        "address": ("address", "მისამართი", "адрес"),
        "acquisition_source": ("source", "how did you hear", "წყარო", "источник"),
        "notes": (
            "notes",
            "note",
            "comment",
            "comments",
            "შენიშვნა",
            "კომენტარი",
            "примечание",
            "комментарий",
        ),
    },
    "translators": {
        "name": (
            "name",
            "full name",
            "translator",
            "თარჯიმანი",
            "სახელი",
            "სახელი გვარი",
            "გვარი სახელი",
            "переводчик",
            "имя",
            "фио",
        ),
        "phone": (
            "phone",
            "tel",
            "telephone",
            "mobile",
            "ტელეფონი",
            "ტელ",
            "მობილური",
            "телефон",
        ),
        "email": ("email", "e-mail", "mail", "ელფოსტა", "ელ ფოსტა", "იმეილი", "почта"),
        "office_address": ("address", "office address", "მისამართი", "адрес"),
        "comment": ("comment", "comments", "notes", "note", "შენიშვნა", "კომენტარი"),
        "default_rate": (
            "rate",
            "price per page",
            "default rate",
            "ტარიფი",
            "ფასი",
            "ფასი გვერდზე",
            "тариф",
        ),
        "bank_iban": ("iban", "account", "bank account", "ანგარიშის ნომერი", "ანგარიში"),
        "can_notarize": (
            "notary",
            "notarize",
            "notarise",
            "can notarize",
            "ნოტარიუსი",
            "ნოტარიული",
            "ნოტარიულად",
            "нотариус",
        ),
        "works_oral": (
            "oral",
            "interpreter",
            "interpreting",
            "ზეპირი",
            "ზეპირი თარგმანი",
            "устный",
        ),
        "cities": ("city", "cities", "ქალაქი", "ქალაქები", "город"),
        "language_pairs": (
            "languages",
            "language",
            "language pairs",
            "pairs",
            "ენები",
            "ენა",
            "მიმართულებები",
            "языки",
        ),
    },
}


FIELDS["notaries"] = (
    FieldSpec("name", required=True),
    FieldSpec("phone"),
    FieldSpec("email"),
    FieldSpec("registration_number"),
    FieldSpec("office_address"),
    FieldSpec("comment"),
    FieldSpec("bank_iban"),
    FieldSpec("bank_inn"),
    FieldSpec("bank_code"),
)

SYNONYMS["notaries"] = {
    "name": (
        "name",
        "full name",
        "notary",
        "ნოტარიუსი",
        "სახელი",
        "სახელი გვარი",
        "გვარი სახელი",
        "нотариус",
        "имя",
        "фио",
    ),
    "phone": (
        "phone",
        "tel",
        "telephone",
        "mobile",
        "ტელეფონი",
        "ტელ",
        "მობილური",
        "телефон",
    ),
    "email": ("email", "e-mail", "mail", "ელფოსტა", "ელ ფოსტა", "იმეილი", "почта"),
    "registration_number": (
        "registration number",
        "registration",
        "reg number",
        "license",
        "licence",
        "license number",
        "სარეგისტრაციო ნომერი",
        "სარეგისტრაციო",
        "რეგისტრაციის ნომერი",
        "ლიცენზია",
        "рег номер",
    ),
    "office_address": ("address", "office address", "მისამართი", "адрес"),
    "comment": ("comment", "comments", "notes", "note", "შენიშვნა", "კომენტარი"),
    "bank_iban": ("iban", "account", "bank account", "ანგარიშის ნომერი", "ანგარიში"),
    "bank_inn": (
        "inn",
        "tax id",
        "tax number",
        "id code",
        "identification code",
        "საიდენტიფიკაციო კოდი",
        "ს კ",
        "პირადი ნომერი",
        "პ ნ",
    ),
    "bank_code": ("bank code", "swift", "bic", "ბანკის კოდი", "ბანკი"),
}

#: The fields that say what a sheet IS. A sheet with a registration number and a
#: bank code is a list of notaries however its name column is called.
DISTINCTIVE: dict[str, frozenset[str]] = {
    "clients": frozenset({"client_type", "personal_number", "address", "acquisition_source"}),
    "translators": frozenset(
        {"default_rate", "can_notarize", "works_oral", "cities", "language_pairs"}
    ),
    "notaries": frozenset({"registration_number", "bank_inn", "bank_code"}),
    "orders": frozenset(
        {"order_date", "price", "pages", "language_pair", "document_type", "paid", "status"}
    ),
}


FIELDS["orders"] = (
    FieldSpec("client_name"),
    FieldSpec("client_phone"),
    FieldSpec("client_email"),
    FieldSpec("client_personal_number"),
    FieldSpec("order_date", required=True),
    FieldSpec("due_date"),
    FieldSpec("language_pair"),
    FieldSpec("source_language"),
    FieldSpec("target_language"),
    FieldSpec("document_type"),
    FieldSpec("pages"),
    FieldSpec("price", required=True),
    FieldSpec("translator"),
    FieldSpec("translator_phone"),
    FieldSpec("translator_cost"),
    FieldSpec("notary"),
    FieldSpec("notary_phone"),
    FieldSpec("notary_cost"),
    FieldSpec("status"),
    FieldSpec("paid"),
    FieldSpec("identifier"),
    FieldSpec("notes"),
)

SYNONYMS["orders"] = {
    "client_name": (
        "client",
        "client name",
        "customer",
        "name",
        "კლიენტი",
        "კლიენტის სახელი",
        "დამკვეთი",
        "სახელი გვარი",
        "клиент",
        "заказчик",
    ),
    "client_phone": (
        "phone",
        "tel",
        "mobile",
        "client phone",
        "ტელეფონი",
        "ტელ",
        "მობილური",
        "телефон",
    ),
    "client_email": ("email", "e-mail", "client email", "ელფოსტა", "იმეილი", "почта"),
    "client_personal_number": (
        "id",
        "id number",
        "personal number",
        "პირადი ნომერი",
        "პირადობის ნომერი",
        "საიდენტიფიკაციო კოდი",
        "ს კ",
        "პ/ნ ან ს/კ",
        "პ ნ",
    ),
    "order_date": (
        "date",
        "order date",
        "created",
        "received",
        "თარიღი",
        "შეკვეთის თარიღი",
        "მიღების თარიღი",
        "дата",
    ),
    "due_date": (
        "due",
        "due date",
        "deadline",
        "ვადა",
        "ბოლო ვადა",
        "გადაცემის თარიღი",
        "срок",
    ),
    "language_pair": (
        "languages",
        "language pair",
        "pair",
        "ენები",
        "ენობრივი წყვილი",
        "მიმართულება",
        "языки",
    ),
    "source_language": ("from", "source", "source language", "საიდან", "ორიგინალის ენა", "с"),
    "target_language": ("to", "target", "target language", "სად", "თარგმანის ენა", "на"),
    "document_type": (
        "document",
        "document type",
        "type",
        "doc type",
        "დოკუმენტი",
        "დოკუმენტის ტიპი",
        "ტიპი",
        "документ",
    ),
    "pages": (
        "pages",
        "page count",
        "qty",
        "quantity",
        "გვერდი",
        "გვერდები",
        "გვერდების რაოდენობა",
        "გვ",
        "страниц",
    ),
    "price": (
        "price",
        "amount",
        "total",
        "cost",
        "sum",
        "ფასი",
        "თანხა",
        "ღირებულება",
        "ჯამი",
        "цена",
        "сумма",
    ),
    "translator": ("translator", "თარჯიმანი", "переводчик"),
    "translator_phone": ("translator phone", "თარჯიმნის ტელეფონი", "тел переводчика"),
    "notary": ("notary", "ნოტარიუსი", "нотариус"),
    "notary_phone": ("notary phone", "ნოტარიუსის ტელეფონი", "тел нотариуса"),
    "notary_cost": (
        "notary cost",
        "notary fee",
        "notary payment",
        "ნოტარიუსის თანხა",
        "ნოტარიუსის ანაზღაურება",
        "ნოტარიუსის საფასური",
    ),
    "translator_cost": (
        "translator cost",
        "translator fee",
        "translator payment",
        "თარჯიმნის ანაზღაურება",
        "თარჯიმნის თანხა",
    ),
    "status": ("status", "სტატუსი", "статус"),
    "paid": ("paid", "paid amount", "payment", "გადახდილი", "გადახდილი თანხა", "оплачено"),
    "identifier": (
        "identifier",
        "tender",
        "tender number",
        "reference",
        "ref",
        "იდენტიფიკატორი",
        "ტენდერი",
        "ტენდერის ნომერი",
    ),
    "notes": ("notes", "note", "comment", "comments", "შენიშვნა", "კომენტარი", "примечание"),
}


def norm_header(text: str) -> str:
    """Lower case, letters and digits only, single spaces."""
    lowered = unicodedata.normalize("NFKC", str(text)).casefold()
    cleaned = re.sub(r"[^\w]+", " ", lowered, flags=re.UNICODE).replace("_", " ")
    return " ".join(cleaned.split())


def suggest_mapping(kind: str, headers: list[str]) -> dict[str, int]:
    """Which column holds which field, guessed from the headers.

    Exact synonym first, then a header that contains one. A column serves one
    field only, and a field takes its first matching column.
    """
    wanted = SYNONYMS[kind]
    normalised = [norm_header(h) for h in headers]
    taken: set[int] = set()
    found: dict[str, int] = {}

    for exact in (True, False):
        for spec in FIELDS[kind]:
            if spec.key in found:
                continue
            words = {norm_header(w) for w in wanted.get(spec.key, ())}
            for index, header in enumerate(normalised):
                if index in taken or not header:
                    continue
                hit = header in words if exact else any(len(w) > 3 and w in header for w in words)
                if hit:
                    found[spec.key] = index
                    taken.add(index)
                    break
    if kind == "orders":
        _assign_phone_columns(normalised, found, SYNONYMS["orders"]["client_phone"])
    return found


_PHONE_OWNERS = (
    ("client_phone", "client_name"),
    ("translator_phone", "translator"),
    ("notary_phone", "notary"),
)


def _assign_phone_columns(
    normalised: list[str], found: dict[str, int], phone_words: tuple[str, ...]
) -> None:
    """Give each "phone" column to the person named just to its left.

    A flat sheet has one phone header three times (the client's, the
    translator's, the notary's). The first match per field would hand them out
    in field order, which is wrong as soon as one person has no phone column.
    Position says it better: a phone column belongs to the nearest name column
    on its left, and to the client when there is none.
    """
    words = {norm_header(w) for w in phone_words}
    phone_columns = [i for i, header in enumerate(normalised) if header in words]
    if not phone_columns:
        return
    for phone_field, _ in _PHONE_OWNERS:
        found.pop(phone_field, None)
    anchors = [
        (found[name_field], phone_field)
        for phone_field, name_field in _PHONE_OWNERS
        if name_field in found
    ]
    for column in phone_columns:
        before = [anchor for anchor in anchors if anchor[0] < column]
        owner = max(before)[1] if before else "client_phone"
        if owner not in found:
            found[owner] = column


# ── Reading the file ────────────────────────────────────────────────────────


@dataclass
class Sheet:
    headers: list[str]
    #: (spreadsheet row number, cells) for each non-empty row after the header.
    rows: list[tuple[int, list[Any]]]


@dataclass
class RawSheet:
    """One worksheet as it was read: its name and its rows that are not blank."""

    name: str
    #: (spreadsheet row number, cells), blank rows left out.
    rows: list[tuple[int, list[Any]]]


_OLE_MAGIC = b"\xd0\xcf\x11\xe0"


def read_workbook(content: bytes, filename: str) -> list[RawSheet]:
    """Every visible sheet of an .xlsx or an old .xls, or a .csv as one sheet.

    Sheets without a single non-blank row are left out. Nothing is guessed here:
    which row is the header, and what a sheet is, is for `imports_combined`.
    """
    if len(content) > MAX_BYTES:
        raise ValidationError(
            f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB. "
            "Split it and import it in parts."
        )
    if not content:
        raise ValidationError("The file is empty.")

    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or zipfile.is_zipfile(io.BytesIO(content)):
        sheets = _read_xlsx(content)
    elif name.endswith(".xls") or content.startswith(_OLE_MAGIC):
        sheets = _read_xls(content)
    elif name.endswith((".csv", ".txt")) or not name.endswith(".ods"):
        stem = (filename or "file").rsplit(".", 1)[0] or "file"
        sheets = [RawSheet(stem, _read_csv(content))]
    else:
        raise ValidationError(
            "Save the file as .xlsx, .xls or .csv first: the .ods format is not read."
        )

    kept = [
        RawSheet(sheet.name, [(n, row) for n, row in sheet.rows if any(cell_text(c) for c in row)])
        for sheet in sheets
    ]
    kept = [sheet for sheet in kept if sheet.rows]
    if not kept:
        raise ValidationError("The file has no rows.")
    return kept


def read_sheet(content: bytes, filename: str) -> Sheet:
    """The first sheet of a workbook, or a .csv, as a header and rows.

    The first row that is not blank is the header. For a file with several
    sheets, or a title above the header, see `imports_combined`.
    """
    (_, header_cells), *body = read_workbook(content, filename)[0].rows
    headers = [cell_text(c) for c in header_cells]
    if len(body) > MAX_ROWS:
        raise ValidationError(
            f"The file has {len(body):,} rows; one import takes at most {MAX_ROWS:,}. Split it."
        )
    if not body:
        raise ValidationError("The file has a header but no rows under it.")
    return Sheet(headers=headers, rows=body)


def _read_xlsx(content: bytes) -> list[RawSheet]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds on a bad file
        raise ValidationError("That file could not be read as an Excel workbook.") from exc
    try:
        return [
            RawSheet(
                sheet.title,
                [
                    (number, list(row))
                    for number, row in enumerate(sheet.iter_rows(values_only=True), start=1)
                ],
            )
            for sheet in workbook.worksheets
            if sheet.sheet_state == "visible"
        ]
    finally:
        workbook.close()


def _read_xls(content: bytes) -> list[RawSheet]:
    """The old Excel 97-2003 format, every visible sheet, dates turned into dates."""
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:  # xlrd raises its own error and others on a damaged file
        raise ValidationError(
            "That file could not be read as an old Excel (.xls) workbook. "
            "Save it as .xlsx and try again."
        ) from exc
    try:
        sheets = []
        for sheet in book.sheets():
            if sheet.visibility != 0:
                continue
            rows = [
                (
                    number + 1,
                    [
                        _xls_value(sheet.cell(number, column), book.datemode)
                        for column in range(sheet.ncols)
                    ],
                )
                for number in range(sheet.nrows)
            ]
            sheets.append(RawSheet(sheet.name, rows))
        return sheets
    finally:
        book.release_resources()


def _xls_value(cell: Any, datemode: int) -> Any:
    """One .xls cell as the plain value openpyxl would have given."""
    import xlrd

    kind = cell.ctype
    if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if kind == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if kind == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        except (xlrd.xldate.XLDateError, OverflowError, ValueError):
            return cell.value
    return cell.value


def _read_csv(content: bytes) -> list[tuple[int, list[Any]]]:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [(number, list(row)) for number, row in enumerate(reader, start=1)]


def _decode(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


# ── Cleaning a cell ─────────────────────────────────────────────────────────


def cell_text(value: Any) -> str:
    """A cell as text. A whole number Excel stored as a float ("555123456.0") is not."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return " ".join(str(value).replace(" ", " ").split())


_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

_COMPANY_WORDS = {
    "b2b",
    "company",
    "organisation",
    "organization",
    "org",
    "legal",
    "business",
    "კომპანია",
    "ორგანიზაცია",
    "იურიდიული",
    "იურიდიული პირი",
    "ბიზნესი",
    "юрлицо",
}
_PERSON_WORDS = {
    "b2c",
    "person",
    "individual",
    "private",
    "physical",
    "ფიზიკური",
    "ფიზიკური პირი",
    "პირი",
}
_COMPANY_MARKS = re.compile(r"\b(llc|ltd|inc|gmbh|jsc|შპს|სს|ააიპ|ი/მ|ип|ооо)\b", re.IGNORECASE)
_YES = {"yes", "y", "true", "1", "x", "+", "კი", "დიახ", "да", "✓", "✔"}
_NO = {"no", "n", "false", "0", "-", "არა", "нет", ""}


def clean_phone(value: str) -> str | None:
    """As typed, spaces tidied; None when nothing that looks like a number is left."""
    digits = re.sub(r"\D", "", value)
    return value[:50] if len(digits) >= 6 else None


def phone_key(value: str | None) -> str | None:
    """The last nine digits, so "+995 555 12-34-56" and "555123456" are one number."""
    digits = re.sub(r"\D", "", value or "")
    return digits[-9:] if len(digits) >= 6 else None


def parse_bool(value: str) -> bool | None:
    """True, False, or None for something that is neither (and so is left unset)."""
    token = value.casefold().strip()
    if token in _YES:
        return True
    if token in _NO:
        return False
    return None


def parse_rate(value: str) -> Decimal | None:
    """ "25,50" and "25.5 GEL" as 25.50; None when there is no amount."""
    cleaned = re.sub(r"[^\d,.\-]", "", value).replace(",", ".")
    if not cleaned or cleaned.count(".") > 1:
        return None
    try:
        amount = Decimal(cleaned)
    except InvalidOperation:
        return None
    return amount if Decimal(0) <= amount <= Decimal(100000) else None


_DATE_FORMATS = (
    "%d.%m.%Y",
    "%d.%m.%y",
    "%d/%m/%Y",
    "%d/%m/%y",
    "%d-%m-%Y",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%d.%m.%Y %H:%M",
    "%Y-%m-%d %H:%M:%S",
)


def parse_date(value: Any) -> date | None:
    """A date from an Excel date cell or from text such as "25.10.2025" or "2025-10-25".

    Day first, as Georgian sheets are written. None for anything else, and for
    years outside 2000 to 2100, which are a typo or a serial number.
    """
    if isinstance(value, datetime):
        found: date | None = value.date()
    elif isinstance(value, date):
        found = value
    else:
        text = cell_text(value)
        found = None
        for pattern in _DATE_FORMATS:
            try:
                found = datetime.strptime(text, pattern).date()  # noqa: DTZ007 (a calendar date, no zone)
                break
            except ValueError:
                continue
    return found if found and 2000 <= found.year <= 2100 else None


def parse_int(value: str) -> int | None:
    """ "3", "3.0" and "3 pages" as 3; None when there is no whole number."""
    match = re.search(r"\d+", value.replace(",", "."))
    return int(match.group()) if match else None


# ── Judging rows ────────────────────────────────────────────────────────────


class RefCounter:
    """Hands out the refs ("new:0", "new:1", ...) of records one import creates.

    One counter per kind of record for the whole import, however many sheets it
    has, so two sheets can never both create "new:0".
    """

    def __init__(self, prefix: str) -> None:
        self.prefix = prefix
        self.count = 0

    def __iter__(self) -> RefCounter:
        return self

    def __next__(self) -> str:
        ref = self.peek()
        self.count += 1
        return ref

    def peek(self) -> str:
        """The ref the next record will get, without using it up."""
        return f"{self.prefix}:{self.count}"


@dataclass
class RowResult:
    row: int
    status: Literal["ok", "duplicate", "error"]
    #: What is shown for the row in the preview: its name, and what will be saved.
    name: str = ""
    values: dict[str, Any] = field(default_factory=dict)
    #: Why it is a duplicate or an error.
    messages: list[str] = field(default_factory=list)
    #: Things left out of an otherwise good row.
    warnings: list[str] = field(default_factory=list)
    #: The existing record a duplicate matched, as "number name".
    matched: str | None = None
    #: What a record this row creates is called until it has an id ("new:3"): set
    #: only when the import needs later rows to point at it.
    ref: str | None = None
    #: For a duplicate: what the record it matched is called elsewhere (an id, or the
    #: ref of one this import creates), so a later row naming this spelling finds it.
    matched_ref: int | str | None = None


@dataclass
class ExistingPerson:
    """One record already in the bureau, for duplicate matching."""

    label: str
    phone: str | None = None
    email: str | None = None
    personal_number: str | None = None
    name: str | None = None
    #: What the record is called elsewhere: a client's id, or a placeholder for one
    #: this import is about to create.
    ref: int | str | None = None


class Seen:
    """Who is already known: the bureau's records, then the rows accepted so far."""

    def __init__(self, existing: Iterable[ExistingPerson] = ()) -> None:
        # Each maps to (label, ref).
        self.by_phone: dict[str, tuple[str, int | str | None]] = {}
        self.by_email: dict[str, tuple[str, int | str | None]] = {}
        self.by_number: dict[str, tuple[str, int | str | None]] = {}
        self.by_name: dict[str, tuple[str, int | str | None]] = {}
        for person in existing:
            self.add(person)

    def add(self, person: ExistingPerson) -> None:
        entry = (person.label, person.ref)
        if key := phone_key(person.phone):
            self.by_phone.setdefault(key, entry)
        if person.email:
            self.by_email.setdefault(person.email.casefold(), entry)
        if person.personal_number:
            self.by_number.setdefault(person.personal_number.casefold(), entry)
        if person.name:
            self.by_name.setdefault(person.name.casefold(), entry)

    def match_ref(
        self, person: ExistingPerson, *, by_name: bool = False
    ) -> tuple[str, str, int | str | None] | None:
        """(what matched, the record's label, its ref), or None."""
        if (key := phone_key(person.phone)) and key in self.by_phone:
            return ("phone", *self.by_phone[key])
        if person.email and person.email.casefold() in self.by_email:
            return ("email", *self.by_email[person.email.casefold()])
        if person.personal_number and person.personal_number.casefold() in self.by_number:
            return ("ID number", *self.by_number[person.personal_number.casefold()])
        if by_name and person.name and person.name.casefold() in self.by_name:
            return ("name", *self.by_name[person.name.casefold()])
        return None

    def match(self, person: ExistingPerson, *, by_name: bool = False) -> tuple[str, str] | None:
        """(what matched, the record's label), or None."""
        hit = self.match_ref(person, by_name=by_name)
        return (hit[0], hit[1]) if hit else None


def _cell(cells: list[Any], mapping: dict[str, int], key: str) -> str:
    index = mapping.get(key)
    if index is None or index >= len(cells):
        return ""
    return cell_text(cells[index])


def _common(
    cells: list[Any], mapping: dict[str, int], result: RowResult
) -> tuple[str, str | None, str | None]:
    """Name, phone and email, cleaned, with what was left out noted."""
    name = _cell(cells, mapping, "name")
    phone_raw = _cell(cells, mapping, "phone")
    email_raw = _cell(cells, mapping, "email").casefold()

    phone = clean_phone(phone_raw) if phone_raw else None
    if phone_raw and phone is None:
        result.warnings.append(f"Phone left out: {phone_raw[:30]!r} is not a number.")
    email: str | None = None
    if email_raw:
        if _EMAIL.match(email_raw) and len(email_raw) <= 255:
            email = email_raw
        else:
            result.warnings.append(f"Email left out: {email_raw[:40]!r} is not valid.")
    return name, phone, email


def judge_client(
    number: int,
    cells: list[Any],
    mapping: dict[str, int],
    seen: Seen,
    refs: Iterator[str] | None = None,
) -> RowResult:
    result = RowResult(row=number, status="ok")
    name, phone, email = _common(cells, mapping, result)
    result.name = name

    if not name:
        result.status = "error"
        result.messages.append("No name.")
        return result
    if len(name) > 255:
        result.status = "error"
        result.messages.append("The name is longer than 255 characters.")
        return result

    kind_text = _cell(cells, mapping, "client_type").casefold()
    if kind_text in _COMPANY_WORDS:
        client_type = "B2B"
    elif kind_text in _PERSON_WORDS:
        client_type = "B2C"
    else:
        client_type = "B2B" if _COMPANY_MARKS.search(name) else "B2C"

    personal = _cell(cells, mapping, "personal_number")[:50] or None
    address = _cell(cells, mapping, "address")[:500] or None
    source = _cell(cells, mapping, "acquisition_source")[:255] or None
    notes = _cell(cells, mapping, "notes") or None

    person = ExistingPerson(
        label=name, phone=phone, email=email, personal_number=personal, name=None
    )
    hit = seen.match_ref(person)
    if hit:
        result.status = "duplicate"
        result.matched = hit[1]
        result.matched_ref = hit[2]
        result.messages.append(f"Already in the list by {hit[0]}: {hit[1]}.")
        return result

    if refs is not None:
        result.ref = person.ref = next(refs)
        # An order later in the file can now name this client, by name too.
        person.name = name
    seen.add(person)
    result.values = {
        "name": name,
        "client_type": client_type,
        "phone": phone,
        "email": email,
        "personal_number": personal,
        "address": address,
        "acquisition_source": source,
        "notes": notes,
    }
    return result


# ── Translators ─────────────────────────────────────────────────────────────

CITY_NAMES: dict[str, tuple[str, str]] = {
    "tbilisi": ("Tbilisi", "თბილისი"),
    "batumi": ("Batumi", "ბათუმი"),
    "kutaisi": ("Kutaisi", "ქუთაისი"),
    "rustavi": ("Rustavi", "რუსთავი"),
    "gori": ("Gori", "გორი"),
    "zugdidi": ("Zugdidi", "ზუგდიდი"),
    "poti": ("Poti", "ფოთი"),
    "telavi": ("Telavi", "თელავი"),
    "akhaltsikhe": ("Akhaltsikhe", "ახალციხე"),
    "ozurgeti": ("Ozurgeti", "ოზურგეთი"),
    "kobuleti": ("Kobuleti", "ქობულეთი"),
    "mtskheta": ("Mtskheta", "მცხეთა"),
    "sighnaghi": ("Sighnaghi", "სიღნაღი"),
    "borjomi": ("Borjomi", "ბორჯომი"),
    "marneuli": ("Marneuli", "მარნეული"),
    "gardabani": ("Gardabani", "გარდაბანი"),
    "samtredia": ("Samtredia", "სამტრედია"),
    "senaki": ("Senaki", "სენაკი"),
    "khashuri": ("Khashuri", "ხაშური"),
    "sachkhere": ("Sachkhere", "საჩხერე"),
}
assert set(CITY_NAMES) == set(CITIES), "CITY_NAMES must list every city in domain.cities"

_SPLIT = re.compile(r"[,;/\n|]+")
_DASH = re.compile(r"\s*(?:->|→|>|–|—|-|=>)\s*")


def parse_cities(value: str) -> tuple[list[str], list[str]]:
    """City codes from "Tbilisi, ბათუმი", and the words that were not a known city."""
    lookup: dict[str, str] = {}
    for code, (en, ka) in CITY_NAMES.items():
        lookup[code] = code
        lookup[en.casefold()] = code
        lookup[ka.casefold()] = code
    codes: list[str] = []
    unknown: list[str] = []
    for part in _SPLIT.split(value):
        word = part.strip().casefold()
        if not word:
            continue
        code = lookup.get(word)
        if code is None:
            unknown.append(part.strip())
        elif code not in codes:
            codes.append(code)
    return codes, unknown


def parse_pairs(
    value: str, language_lookup: Callable[[str], str | None]
) -> tuple[list[tuple[str, str]], list[str]]:
    """Language pairs from "ka-en; English > Georgian", and the pieces not understood.

    ``language_lookup`` turns a code or a name (in any language) into the
    bureau's language code, or None.
    """
    pairs: list[tuple[str, str]] = []
    problems: list[str] = []
    for part in _SPLIT.split(value):
        text = part.strip()
        if not text:
            continue
        sides = _DASH.split(text)
        if len(sides) != 2:
            problems.append(text)
            continue
        source, target = (language_lookup(sides[0].strip()), language_lookup(sides[1].strip()))
        if source is None or target is None or source == target:
            problems.append(text)
            continue
        if (source, target) not in pairs:
            pairs.append((source, target))
    return pairs, problems


def judge_translator(
    number: int,
    cells: list[Any],
    mapping: dict[str, int],
    seen: Seen,
    language_lookup: Callable[[str], str | None],
    refs: Iterator[str] | None = None,
) -> RowResult:
    result = RowResult(row=number, status="ok")
    name, phone, email = _common(cells, mapping, result)
    result.name = name

    if not name:
        result.status = "error"
        result.messages.append("No name.")
        return result
    if len(name) > 255:
        result.status = "error"
        result.messages.append("The name is longer than 255 characters.")
        return result

    rate_text = _cell(cells, mapping, "default_rate")
    rate = parse_rate(rate_text) if rate_text else None
    if rate_text and rate is None:
        result.warnings.append(f"Rate left out: {rate_text[:20]!r} is not an amount.")

    notarize = parse_bool(_cell(cells, mapping, "can_notarize"))
    oral = parse_bool(_cell(cells, mapping, "works_oral"))

    cities, unknown_cities = parse_cities(_cell(cells, mapping, "cities"))
    if unknown_cities:
        result.warnings.append("Cities not recognised: " + ", ".join(unknown_cities[:5]) + ".")

    pairs, pair_problems = parse_pairs(_cell(cells, mapping, "language_pairs"), language_lookup)
    if pair_problems:
        result.warnings.append(
            "Language pairs not understood (use ka-en, or English-Georgian): "
            + "; ".join(pair_problems[:5])
            + "."
        )

    person = ExistingPerson(label=name, phone=phone, email=email, name=name)
    hit = seen.match_ref(person, by_name=True)
    if hit:
        result.status = "duplicate"
        result.matched = hit[1]
        result.matched_ref = hit[2]
        result.messages.append(f"Already in the list by {hit[0]}: {hit[1]}.")
        return result

    if refs is not None:
        result.ref = person.ref = next(refs)
    seen.add(person)
    result.values = {
        "name": name,
        "phone": phone,
        "email": email,
        "office_address": _cell(cells, mapping, "office_address")[:255] or None,
        "comment": _cell(cells, mapping, "comment") or None,
        "default_rate": rate,
        "bank_iban": _cell(cells, mapping, "bank_iban").replace(" ", "")[:34] or None,
        "can_notarize": bool(notarize),
        "works_oral": bool(oral),
        "cities": cities,
        "language_pairs": pairs,
    }
    return result


def judge_notary(
    number: int,
    cells: list[Any],
    mapping: dict[str, int],
    seen: Seen,
    refs: Iterator[str] | None = None,
) -> RowResult:
    """A notary: a name, and whatever else the sheet has (no IBAN is fine)."""
    result = RowResult(row=number, status="ok")
    name, phone, email = _common(cells, mapping, result)
    result.name = name

    if not name:
        result.status = "error"
        result.messages.append("No name.")
        return result
    if len(name) > 255:
        result.status = "error"
        result.messages.append("The name is longer than 255 characters.")
        return result

    person = ExistingPerson(label=name, phone=phone, email=email, name=name)
    hit = seen.match_ref(person, by_name=True)
    if hit:
        result.status = "duplicate"
        result.matched = hit[1]
        result.matched_ref = hit[2]
        result.messages.append(f"Already in the list by {hit[0]}: {hit[1]}.")
        return result

    if refs is not None:
        result.ref = person.ref = next(refs)
    seen.add(person)
    result.values = {
        "name": name,
        "phone": phone,
        "email": email,
        "registration_number": _cell(cells, mapping, "registration_number")[:100] or None,
        "office_address": _cell(cells, mapping, "office_address")[:255] or None,
        "comment": _cell(cells, mapping, "comment") or None,
        "bank_iban": _cell(cells, mapping, "bank_iban").replace(" ", "").upper()[:34] or None,
        "bank_inn": _cell(cells, mapping, "bank_inn")[:20] or None,
        "bank_code": _cell(cells, mapping, "bank_code").upper()[:20] or None,
    }
    return result


# ── The whole file ──────────────────────────────────────────────────────────


def check_mapping(kind: str, mapping: dict[str, int], column_count: int) -> dict[str, int]:
    """The mapping, or a ValidationError saying what is wrong with it."""
    allowed = {spec.key for spec in FIELDS[kind]}
    clean: dict[str, int] = {}
    for key, index in mapping.items():
        if key not in allowed:
            raise ValidationError(f"Unknown field: {key}.")
        if not isinstance(index, int) or isinstance(index, bool) or not 0 <= index < column_count:
            raise ValidationError(f"Column for {key} is not in the file.")
        clean[key] = index
    for spec in FIELDS[kind]:
        if spec.required and spec.key not in clean:
            raise ValidationError(f"Choose the column for {spec.key.replace('_', ' ')}.")
    if len(set(clean.values())) != len(clean):
        raise ValidationError("A column can be used for one field only.")
    return clean


@dataclass
class Judged:
    results: list[RowResult]

    def count(self, status: str) -> int:
        return sum(1 for r in self.results if r.status == status)


def judge_sheet(
    kind: str,
    sheet: Sheet,
    mapping: dict[str, int],
    existing: Iterable[ExistingPerson],
    language_lookup: Callable[[str], str | None] | None = None,
    orders: Any = None,
) -> Judged:
    seen = Seen(existing)
    results: list[RowResult] = []
    if kind == "orders":
        from suliko.domain.imports_orders import judge_orders

        return Judged(judge_orders(sheet, mapping, seen, orders))
    for number, cells in sheet.rows:
        if kind == "clients":
            results.append(judge_client(number, cells, mapping, seen))
        elif kind == "notaries":
            results.append(judge_notary(number, cells, mapping, seen))
        else:
            results.append(
                judge_translator(number, cells, mapping, seen, language_lookup or (lambda _t: None))
            )
    return Judged(results)
